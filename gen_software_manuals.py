#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""为三个内置软件生成各自的使用说明书（.docx）"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "manuals")
os.makedirs(OUT_DIR, exist_ok=True)

NAVY = RGBColor(0x2C, 0x3E, 0x6B)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color in [("Heading 1", 17, NAVY), ("Heading 2", 13.5, RGBColor(0x6B, 0x8F, 0xA3))]:
        s = doc.styles[name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(size * 0.75)
        s.paragraph_format.space_after = Pt(size * 0.4)
        s.paragraph_format.keep_with_next = True
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_header_footer(doc, title):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = title
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    ph = OxmlElement("w:t"); ph.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    run._r.extend([f1, it, f2, ph, f3])
    for r in fp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    tblPr = t._tbl.tblPr
    tb = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "C8D6E0")
        tb.append(b)
    tblPr.append(tb)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "6B8FA31A")
        tcPr.append(shd)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0]; r.bold = True; r.font.color.rgb = NAVY
    for ri, row in enumerate(rows):
        for ci, txt in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            c.text = str(txt)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Inches(w)
    doc.add_paragraph()


def build(doc, meta, blocks):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1.1)
    configure_styles(doc)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(160)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["title"])
    r.bold = True; r.font.size = Pt(34); r.font.color.rgb = NAVY
    rPr = r._r.get_or_add_rPr(); rPr.get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(10)
    r2 = p2.add_run("使用说明书")
    r2.bold = True; r2.font.size = Pt(28); r2.font.color.rgb = NAVY
    rPr = r2._r.get_or_add_rPr(); rPr.get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(14)
    r3 = p3.add_run(meta["sub"])
    r3.font.size = Pt(13); r3.font.color.rgb = RGBColor(0x6B, 0x8F, 0xA3)
    p4 = doc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(220)
    r4 = p4.add_run(f"版本：{meta['ver']}（2026 年 9 月）\n配套安装包：{meta['pkg']}")
    r4.font.size = Pt(10.5); r4.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_page_break()

    add_header_footer(doc, meta["title"] + " · 使用说明书")

    for kind, *args in blocks:
        if kind == "h1":
            doc.add_heading(args[0], level=1)
        elif kind == "h2":
            doc.add_heading(args[0], level=2)
        elif kind == "p":
            doc.add_paragraph(args[0])
        elif kind == "num":
            for i, t in enumerate(args[0], 1):
                doc.add_paragraph(f"{i}. {t}", style="List Number")
        elif kind == "bullet":
            for t in args[0]:
                doc.add_paragraph(t, style="List Bullet")
        elif kind == "table":
            add_table(doc, args[0], args[1], args[2] if len(args) > 2 else None)
    doc.save(os.path.join(OUT_DIR, meta["file"]))
    print("saved:", meta["file"])


# ============ 1. AVTools 音视频工具箱 ============
avtools_blocks = [
    ("h1", "一、软件简介"),
    ("p", "AVTools 音视频工具箱是一款一站式音视频处理软件，将格式转换、视频剪辑、音轨提取、批量处理等常用功能整合在一个界面中。无需专业剪辑经验，即可完成日常的音视频处理任务，适合自媒体创作者、办公人员及家庭用户使用。"),
    ("p", "本说明书对应安装包：AVTools_Setup_1.0.0.exe（版本 1.0.0）。"),

    ("h1", "二、安装与启动"),
    ("num", ["双击 AVTools_Setup_1.0.0.exe 启动安装程序。", "按安装向导提示选择安装目录（默认 C:\\Program Files\\AVTools），点击「下一步」。", "等待安装进度完成，点击「完成」。", "双击桌面「AVTools 音视频工具箱」图标，或从开始菜单启动。"]),
    ("table", ["项目", "要求"], [["操作系统", "Windows 7 / 8 / 10 / 11（64 位）"], ["内存", "建议 4GB 及以上"], ["硬盘空间", "安装需 200MB，处理大文件建议预留 10GB"]], [1.6, 4.2]),

    ("h1", "三、功能详解"),
    ("h2", "3.1 格式转换"),
    ("p", "支持 MP4、MKV、AVI、MOV、FLV、WMV 等常见视频格式以及 MP3、WAV、FLAC、AAC 等音频格式的相互转换。"),
    ("num", ["点击主界面的「格式转换」。", "点击「添加文件」或将文件直接拖入窗口。", "在输出设置中选择目标格式，可自定义分辨率、码率、帧率。", "点击「开始转换」，等待进度条完成。"]),
    ("h2", "3.2 视频剪辑"),
    ("p", "提供裁剪、拼接、变速、添加字幕等基础剪辑功能。"),
    ("num", ["选择「视频剪辑」，导入视频文件。", "拖动时间轴上的手柄，选择需要保留的片段。", "如需要拼接多个片段，可继续添加片段并调整顺序。", "点击「导出」，选择输出格式后保存。"]),
    ("h2", "3.3 音轨提取"),
    ("p", "可一键从视频中提取音频，输出为 MP3、WAV 或 FLAC 格式。"),
    ("num", ["选择「音轨提取」，导入视频文件。", "选择输出音频格式（推荐 MP3，通用性好）。", "点击「开始提取」，完成后音频文件保存在输出目录。"]),
    ("h2", "3.4 批量处理"),
    ("p", "支持一次添加多个文件进行批量转换或批量提取，多线程并行处理，大幅节省时间。批量操作时请保持电脑通电并勿关闭程序。"),

    ("h1", "四、操作示例"),
    ("h2", "示例：将 MKV 视频转换为 MP4"),
    ("num", ["打开「格式转换」，将 MKV 文件拖入窗口。", "输出格式选择 MP4，分辨率保持「与原视频一致」。", "码率选择「自动」。", "点击「开始转换」，等待完成。", "在输出目录中找到同名 MP4 文件，即可正常播放。"]),

    ("h1", "五、常见问题"),
    ("table", ["问题", "解答"], [
        ["转换后画质变差怎么办？", "在输出设置中调高码率（建议 8Mbps 以上），或选择「高质量」预设。"],
        ["支持 4K 视频吗？", "支持，4K 转换建议使用硬件加速并预留充足磁盘空间。"],
        ["转换到一半卡住？", "多为文件损坏或磁盘空间不足，请检查源文件完整性及输出盘剩余空间。"],
        ["可以批量添加不同格式吗？", "可以，批量处理支持混合格式，统一转换为所选的目标格式。"],
    ], [2.0, 4.0]),

    ("h1", "六、注意事项"),
    ("bullet", ["转换/提取过程中请勿关闭电脑电源或强制退出程序。", "输出路径建议使用英文或简短路径，避免个别设备兼容问题。", "处理版权视频请遵守相关法律法规。", "本软件不附带任何广告或捆绑安装，请放心使用。"]),
]

# ============ 2. 极致压缩工具 ============
ct_blocks = [
    ("h1", "一、软件简介"),
    ("p", "极致压缩工具是一款高性能文件压缩/解压软件，主打高压缩率与极速解压，支持 ZIP、7Z、TAR 等多种格式，内置智能压缩策略，可对文本、图片、视频等不同类型文件自动选择最优压缩算法。"),
    ("p", "本说明书对应安装包：CompressTool_Setup_1.0.0.exe（版本 1.0.0）。"),

    ("h1", "二、安装与启动"),
    ("num", ["双击 CompressTool_Setup_1.0.0.exe 启动安装程序。", "按向导选择安装目录，点击「下一步」完成安装。", "安装完成后双击桌面图标启动。"]),
    ("table", ["项目", "要求"], [["操作系统", "Windows 7 / 8 / 10 / 11"], ["内存", "建议 2GB 及以上"], ["硬盘空间", "安装需 100MB"]], [1.6, 4.2]),

    ("h1", "三、快速上手"),
    ("h2", "3.1 压缩文件"),
    ("num", ["启动软件，点击「压缩」或直接将要压缩的文件/文件夹拖入窗口。", "选择压缩格式（推荐 7Z，压缩率最高；ZIP 兼容性最好）。", "设置压缩级别（1-9，级别越高体积越小、耗时越长）。", "点击「开始压缩」，完成后在输出目录生成压缩包。"]),
    ("h2", "3.2 解压文件"),
    ("num", ["点击「解压」，选择要解压的压缩包（支持拖入）。", "选择解压目标目录。", "点击「开始解压」，如压缩包有密码会提示输入。"]),
    ("h2", "3.3 右键快捷菜单"),
    ("p", "安装时可选择「集成右键菜单」。启用后，在资源管理器中右键文件即可看到「压缩到…」和「解压到…」选项，无需打开软件即可快速操作。"),

    ("h1", "四、高级功能"),
    ("h2", "4.1 分卷压缩"),
    ("p", "大文件可拆分为多个分卷（如 1GB 一个），便于存储和传输。在压缩设置中选择「分卷大小」即可。"),
    ("h2", "4.2 加密压缩"),
    ("p", "在压缩设置中填写密码即可加密，解压时需要密码。请务必牢记密码，密码丢失无法找回。"),
    ("h2", "4.3 批量压缩"),
    ("p", "支持一次选择多个文件/文件夹批量压缩，每个文件单独生成压缩包，或合并为一个压缩包。"),

    ("h1", "五、操作示例"),
    ("h2", "示例：把项目文件夹压缩成 7Z 并加密"),
    ("num", ["将文件夹拖入压缩窗口。", "格式选择 7Z，压缩级别选择 5（平衡）。", "勾选「加密」，输入密码并确认。", "点击「开始压缩」。", "将生成的 .7z 文件拷贝到目标位置，解压时输入密码即可。"]),

    ("h1", "六、常见问题"),
    ("table", ["问题", "解答"], [
        ["ZIP 和 7Z 选哪个？", "追求压缩率选 7Z；需要发给别人且对方系统未装解压软件时选 ZIP。"],
        ["压缩包密码忘了？", "密码无法找回，只能重新压缩。重要文件建议设置「密码提示」或使用密码管理器。"],
        ["解压提示文件损坏？", "重新下载/拷贝压缩包，或使用「修复压缩包」功能尝试修复。"],
        ["压缩大文件很慢？", "可降低压缩级别（如 1-3），速度会明显提升；也可勾选「多线程」。"],
    ], [2.0, 4.0]),

    ("h1", "七、注意事项"),
    ("bullet", ["压缩加密文件请定期备份密码。", "分卷压缩的分卷必须放在同一目录下才能正常解压。", "压缩超大文件（>10GB）请确保磁盘有足够空间，输出目录与源目录建议不在同一磁盘。"]),
]

# ============ 3. 游戏 Mod AI 翻译器 ============
gmt_blocks = [
    ("h1", "一、软件简介"),
    ("p", "游戏 Mod AI 翻译器是一款利用 AI 大模型为游戏 Mod 提供高质量翻译的工具。它可以自动识别 Mod 中的文本资源，调用 AI 引擎进行语境化翻译，并一键导出为 Mod 可识别的格式，让海量外文 Mod 无障碍游玩。"),
    ("p", "本说明书对应安装包：GameModTranslator-Setup-1.0.exe（版本 1.0）。"),

    ("h1", "二、安装与启动"),
    ("num", ["双击 GameModTranslator-Setup-1.0.exe 启动安装程序。", "按向导选择安装目录，点击「下一步」完成安装。", "安装完成后双击桌面图标启动。首次启动建议在「设置」中配置 AI 引擎参数。"]),
    ("table", ["项目", "要求"], [["操作系统", "Windows 10 / 11（64 位）"], ["内存", "建议 8GB 及以上"], ["网络", "调用 AI 引擎需要联网"], ["硬盘空间", "安装需 300MB，Mod 缓存建议预留 5GB"]], [1.6, 4.2]),

    ("h1", "三、快速上手"),
    ("h2", "3.1 翻译一个 Mod"),
    ("num", ["点击「添加 Mod」，选择需要翻译的 Mod 文件或文件夹（支持 .pak、.zip、.json 等常见格式）。", "软件自动扫描并列出可翻译的文本条目。", "选择源语言与目标语言（如 英文 → 简体中文）。", "点击「开始翻译」，AI 引擎自动逐条翻译。", "翻译完成后点击「导出」，生成翻译后的 Mod 文件。", "将导出的文件替换/安装到游戏 Mod 目录，重新启动游戏即可生效。"]),
    ("h2", "3.2 术语与风格控制"),
    ("p", "在「术语表」中可以预设专有名词的翻译（如角色名、地名），在「风格」中可设置翻译语气（正式/口语/游戏化），让翻译更贴合游戏语境。"),

    ("h1", "四、功能详解"),
    ("h2", "4.1 批量翻译"),
    ("p", "支持一次添加多个 Mod 排队翻译，适合大型 Mod 合集。批量翻译会消耗较多 AI 配额，建议按需设置每日翻译上限。"),
    ("h2", "4.2 原文对比"),
    ("p", "翻译界面支持原文/译文左右对照查看，可手动微调任意一条翻译结果，修改后自动保存。"),
    ("h2", "4.3 更新检测"),
    ("p", "Mod 有更新时，可导入新版 Mod 文件，软件会识别新增/改动的文本，只翻译差异部分，节省配额和时间。"),

    ("h1", "五、操作示例"),
    ("h2", "示例：翻译并安装一个英文 Mod"),
    ("num", ["准备 Mod 文件（假设为 xxx.pak）。", "打开游戏 Mod AI 翻译器，点击「添加 Mod」，选择 xxx.pak。", "语言设置：英文 → 简体中文，风格选「游戏化」。", "点击「开始翻译」，等待进度完成。", "点击「导出」，保存翻译后的文件。", "将导出文件放入游戏 Mods 目录，启动游戏验证。"]),

    ("h1", "六、常见问题"),
    ("table", ["问题", "解答"], [
        ["提示 AI 配额不足？", "在「设置」中检查 AI 引擎的 API Key 与配额，或切换备用模型。"],
        ["翻译结果不理想？", "可在「术语表」补充专有名词，或手动修正个别条目后重新导出。"],
        ["导出的 Mod 游戏内不生效？", "确认导出格式与游戏 Mod 版本兼容，并检查文件是否放入正确的 Mod 目录。"],
        ["支持哪些语言？", "支持中、英、日、韩、德、法、俄、西等主流语言互译，具体以 AI 引擎支持为准。"],
    ], [2.0, 4.0]),

    ("h1", "七、注意事项"),
    ("bullet", ["使用 AI 翻译需要联网，请确保网络稳定。", "请遵守 Mod 作者的许可协议，勿将翻译结果用于商业用途。", "翻译大型 Mod 请耐心等待，中途退出会中断进度。", "本软件不附带任何广告或捆绑安装，请放心使用。"]),
]

if __name__ == "__main__":
    from docx import Document
    metas = [
        {"title": "AVTools 音视频工具箱", "sub": "一站式音视频处理工具", "ver": "1.0.0", "pkg": "AVTools_Setup_1.0.0.exe", "file": "AVTools_使用说明书.docx"},
        {"title": "极致压缩工具", "sub": "高性能压缩 / 解压工具", "ver": "1.0.0", "pkg": "CompressTool_Setup_1.0.0.exe", "file": "极致压缩工具_使用说明书.docx"},
        {"title": "游戏 Mod AI 翻译器", "sub": "AI 驱动的 Mod 翻译工具", "ver": "1.0", "pkg": "GameModTranslator-Setup-1.0.exe", "file": "游戏ModAI翻译器_使用说明书.docx"},
    ]
    for meta, blocks in zip(metas, [avtools_blocks, ct_blocks, gmt_blocks]):
        build(Document(), meta, blocks)
